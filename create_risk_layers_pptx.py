#!/usr/bin/env python3
"""Generate PowerPoint: Trading System Risk Layers sketch."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Palette
NAVY = RGBColor(0x1A, 0x2B, 0x4A)
DARK = RGBColor(0x2D, 0x37, 0x48)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xF7, 0xFA, 0xFC)
MID_GRAY = RGBColor(0x71, 0x85, 0x9A)
ACCENT = RGBColor(0x31, 0x82, 0xCE)
RISK_RED = RGBColor(0xC5, 0x30, 0x30)
RISK_BG = RGBColor(0xFE, 0xF2, 0xF2)
LAYER_COLORS = [
    RGBColor(0x31, 0x82, 0xCE),  # Trading - blue
    RGBColor(0x80, 0x5A, 0xD5),  # Intermediation - purple
    RGBColor(0xD9, 0x77, 0x06),  # Clearing - amber
    RGBColor(0x05, 0x96, 0x69),  # Custody - green
]
INSIGHT_BG = RGBColor(0xEB, 0xF8, 0xFF)


def set_slide_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_header_bar(slide, title, subtitle=None, accent=LAYER_COLORS[0]):
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.15)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.18), Inches(12.3), Inches(0.55))
    p = tb.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = WHITE

    if subtitle:
        sb = slide.shapes.add_textbox(Inches(0.5), Inches(0.68), Inches(12.3), Inches(0.4))
        sp = sb.text_frame.paragraphs[0]
        sp.text = subtitle
        sp.font.size = Pt(14)
        sp.font.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)


def add_bullet_block(slide, left, top, width, height, items, font_size=14, title=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP

    if title:
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(font_size + 2)
        p.font.bold = True
        p.font.color.rgb = DARK
        p.space_after = Pt(8)

    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 and not title else tf.add_paragraph()
        if title or i > 0:
            p = tf.add_paragraph() if title and i == 0 else p
        # Fix paragraph indexing
        if title:
            p = tf.add_paragraph()
        elif i > 0:
            p = tf.add_paragraph()
        else:
            p = tf.paragraphs[0]

        text, bold, color = item if isinstance(item, tuple) else (item, False, DARK)
        p.text = f"• {text}" if not text.startswith("•") and not text.startswith("→") and not text.startswith("👉") else text
        p.font.size = Pt(font_size)
        p.font.bold = bold
        p.font.color.rgb = color
        p.space_after = Pt(6)
        p.level = 0


def add_risk_box(slide, left, top, width, height, risks):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RISK_BG
    shape.line.color.rgb = RISK_RED
    shape.line.width = Pt(1.5)

    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.15)
    tf.margin_right = Inches(0.15)
    tf.margin_top = Inches(0.12)

    p = tf.paragraphs[0]
    p.text = "ΡΙΣΚΑ"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = RISK_RED
    p.space_after = Pt(6)

    for risk in risks:
        p = tf.add_paragraph()
        p.text = f"• {risk}"
        p.font.size = Pt(12)
        p.font.color.rgb = DARK
        p.space_after = Pt(4)


def add_note_box(slide, left, top, width, text):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, Inches(0.55)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = INSIGHT_BG
    shape.line.color.rgb = ACCENT
    shape.line.width = Pt(1)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.12)
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(12)
    p.font.italic = True
    p.font.color.rgb = NAVY


def slide_title(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, NAVY)

    # Accent stripe
    stripe = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(3.2), Inches(13.333), Inches(0.08)
    )
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = ACCENT
    stripe.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(1.2))
    p = tb.text_frame.paragraphs[0]
    p.text = "Σκαρίφημα Συστήματος Συναλλαγών"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.LEFT

    sub = slide.shapes.add_textbox(Inches(0.8), Inches(2.5), Inches(11.7), Inches(0.6))
    sp = sub.text_frame.paragraphs[0]
    sp.text = "Πού «ζει» το ρίσκο σε κάθε επίπεδο"
    sp.font.size = Pt(22)
    sp.font.color.rgb = RGBColor(0xA0, 0xAE, 0xC0)

    layers = slide.shapes.add_textbox(Inches(0.8), Inches(3.6), Inches(11.7), Inches(0.5))
    lp = layers.text_frame.paragraphs[0]
    lp.text = "Trading  →  Intermediation  →  Clearing/Settlement  →  Custody"
    lp.font.size = Pt(16)
    lp.font.color.rgb = ACCENT


def slide_overview(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, LIGHT_GRAY)
    add_header_bar(slide, "Τα 4 Επίπεδα — Επισκόπηση", "Από το match μέχρι την τελική αποθήκευση assets")

    layers = [
        ("1. TRADING", "Tradeweb κ.λπ.", "Match & Confirmation", "Market risk · Execution risk", "❌ Χωρίς transfer χρημάτων/τίτλων"),
        ("2. INTERMEDIATION", "Tradition κ.λπ.", "Broker / Matched Principal", "Broker credit · Operational risk", "Νομική μεσολάβηση στο trade"),
        ("3. CLEARING / SETTLEMENT", "DVP / CCP", "Cash ↔ Securities", "Principal (↓) · CCP default · Settlement fail", "Το κρίσιμο σημείο"),
        ("4. CUSTODY", "Custodian Bank", "Segregated / Omnibus", "Insolvency · Legal ownership", "Τελική «αποθήκη» assets"),
    ]

    y = 1.45
    for i, (name, example, action, risks, note) in enumerate(layers):
        color = LAYER_COLORS[i]
        # Layer block
        rect = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(y), Inches(12.3), Inches(1.15)
        )
        rect.fill.solid()
        rect.fill.fore_color.rgb = WHITE
        rect.line.color.rgb = color
        rect.line.width = Pt(2.5)

        # Color tag
        tag = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(y), Inches(0.12), Inches(1.15)
        )
        tag.fill.solid()
        tag.fill.fore_color.rgb = color
        tag.line.fill.background()

        # Layer name
        nb = slide.shapes.add_textbox(Inches(0.75), Inches(y + 0.08), Inches(3.2), Inches(0.35))
        np = nb.text_frame.paragraphs[0]
        np.text = name
        np.font.size = Pt(14)
        np.font.bold = True
        np.font.color.rgb = color

        ex = slide.shapes.add_textbox(Inches(0.75), Inches(y + 0.38), Inches(3.2), Inches(0.3))
        ep = ex.text_frame.paragraphs[0]
        ep.text = example
        ep.font.size = Pt(11)
        ep.font.color.rgb = MID_GRAY

        act = slide.shapes.add_textbox(Inches(3.9), Inches(y + 0.15), Inches(3.0), Inches(0.8))
        ap = act.text_frame.paragraphs[0]
        ap.text = action
        ap.font.size = Pt(12)
        ap.font.color.rgb = DARK

        rsk = slide.shapes.add_textbox(Inches(7.0), Inches(y + 0.15), Inches(3.3), Inches(0.8))
        rp = rsk.text_frame.paragraphs[0]
        rp.text = f"⚠ {risks}"
        rp.font.size = Pt(11)
        rp.font.color.rgb = RISK_RED

        nt = slide.shapes.add_textbox(Inches(10.2), Inches(y + 0.15), Inches(2.4), Inches(0.8))
        ntp = nt.text_frame.paragraphs[0]
        ntp.text = note
        ntp.font.size = Pt(10)
        ntp.font.italic = True
        ntp.font.color.rgb = MID_GRAY

        # Arrow between layers
        if i < 3:
            arr = slide.shapes.add_shape(
                MSO_SHAPE.DOWN_ARROW, Inches(6.4), Inches(y + 1.12), Inches(0.35), Inches(0.28)
            )
            arr.fill.solid()
            arr.fill.fore_color.rgb = MID_GRAY
            arr.line.fill.background()

        y += 1.42


def slide_layer(prs, num, title, example, what_happens, risks, note, extra_sections=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, LIGHT_GRAY)
    color = LAYER_COLORS[num - 1]
    add_header_bar(slide, f"Layer {num}: {title}", example, color)

    # What happens
    wh_shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.4), Inches(6.0), Inches(2.8)
    )
    wh_shape.fill.solid()
    wh_shape.fill.fore_color.rgb = WHITE
    wh_shape.line.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)

    tf = wh_shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_top = Inches(0.15)
    p = tf.paragraphs[0]
    p.text = "Τι συμβαίνει"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = color
    p.space_after = Pt(10)

    for item in what_happens:
        p = tf.add_paragraph()
        p.text = f"• {item}"
        p.font.size = Pt(13)
        p.font.color.rgb = DARK
        p.space_after = Pt(6)

    add_risk_box(slide, Inches(6.8), Inches(1.4), Inches(5.9), Inches(2.8), risks)

    y_extra = 4.4
    if extra_sections:
        for sec_title, sec_items in extra_sections:
            box = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(y_extra), Inches(12.2), Inches(0.15 + 0.35 * len(sec_items))
            )
            box.fill.solid()
            box.fill.fore_color.rgb = WHITE
            box.line.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
            stf = box.text_frame
            stf.word_wrap = True
            stf.margin_left = Inches(0.2)
            stf.margin_top = Inches(0.1)
            sp = stf.paragraphs[0]
            sp.text = sec_title
            sp.font.size = Pt(13)
            sp.font.bold = True
            sp.font.color.rgb = DARK
            sp.space_after = Pt(4)
            for it in sec_items:
                ip = stf.add_paragraph()
                ip.text = f"  → {it}"
                ip.font.size = Pt(12)
                ip.font.color.rgb = DARK
                ip.space_after = Pt(3)
            y_extra += 0.2 + 0.35 * len(sec_items)

    if note:
        add_note_box(slide, Inches(0.5), Inches(6.85 if not extra_sections else min(6.85, y_extra + 0.3)), Inches(12.2), note)


def slide_flow(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, LIGHT_GRAY)
    add_header_bar(slide, "FULL FLOW — Mental Model", "Πώς ρέει ένα trade από match μέχρι custody")

    flow_text = """Trader A                    Trader B
    |                            |
    +-------- Tradeweb / Broker (Tradition) --------+
    |              match / confirmation              |
    +------------------------+-----------------------+
                             |
                             v
                  Trade enriched & confirmed
                             |
                             v
              CCP or bilateral DVP settlement
              (cash <-> securities exchange)
                             |
                             v
                    Custodian accounts"""

    box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.5), Inches(10.3), Inches(3.8)
    )
    box.fill.solid()
    box.fill.fore_color.rgb = WHITE
    box.line.color.rgb = NAVY
    box.line.width = Pt(2)

    tb = slide.shapes.add_textbox(Inches(1.8), Inches(1.75), Inches(9.7), Inches(3.4))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = flow_text
    p.font.size = Pt(15)
    p.font.name = "Courier New"
    p.font.color.rgb = DARK

    # Risk timeline
    timeline = [
        ("Πριν το settlement", "Market + counterparty expectation risk", RGBColor(0x31, 0x82, 0xCE)),
        ("Κατά το settlement", "Operational + liquidity + system risk", RGBColor(0xD9, 0x77, 0x06)),
        ("Μετά το settlement", "Custody + legal ownership risk", RGBColor(0x05, 0x96, 0x69)),
    ]
    x = 0.8
    for label, risk, col in timeline:
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(5.6), Inches(3.8), Inches(1.2)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = col
        card.line.width = Pt(2)
        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = Inches(0.15)
        ctf.margin_top = Inches(0.12)
        cp = ctf.paragraphs[0]
        cp.text = label
        cp.font.size = Pt(13)
        cp.font.bold = True
        cp.font.color.rgb = col
        cp.space_after = Pt(6)
        rp = ctf.add_paragraph()
        rp.text = f"→ {risk}"
        rp.font.size = Pt(11)
        rp.font.color.rgb = DARK
        x += 4.1


def slide_insight(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, LIGHT_GRAY)
    add_header_bar(slide, "Το Κρίσιμο Insight", accent=NAVY)

    # DVP does NOT mean
    left = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.6), Inches(5.8), Inches(2.2)
    )
    left.fill.solid()
    left.fill.fore_color.rgb = RISK_BG
    left.line.color.rgb = RISK_RED
    left.line.width = Pt(2)
    ltf = left.text_frame
    ltf.word_wrap = True
    ltf.margin_left = Inches(0.2)
    ltf.margin_top = Inches(0.2)
    lp = ltf.paragraphs[0]
    lp.text = "DVP ΔΕΝ σημαίνει:"
    lp.font.size = Pt(18)
    lp.font.bold = True
    lp.font.color.rgb = RISK_RED
    lp.space_after = Pt(12)
    lp2 = ltf.add_paragraph()
    lp2.text = '"Δεν υπάρχει ρίσκο"'
    lp2.font.size = Pt(22)
    lp2.font.italic = True
    lp2.font.color.rgb = DARK

    right = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.6), Inches(5.8), Inches(2.2)
    )
    right.fill.solid()
    right.fill.fore_color.rgb = RGBColor(0xF0, 0xFF, 0xF4)
    right.line.color.rgb = RGBColor(0x05, 0x96, 0x69)
    right.line.width = Pt(2)
    rtf = right.text_frame
    rtf.word_wrap = True
    rtf.margin_left = Inches(0.2)
    rtf.margin_top = Inches(0.2)
    rp = rtf.paragraphs[0]
    rp.text = "DVP ΣΗΜΑΙΝΕΙ:"
    rp.font.size = Pt(18)
    rp.font.bold = True
    rp.font.color.rgb = RGBColor(0x05, 0x96, 0x69)
    rp.space_after = Pt(12)
    rp2 = rtf.add_paragraph()
    rp2.text = "Το πιο επικίνδυνο ρίσκο\n(loss of principal από exchange failure)\nέχει σχεδόν μηδενιστεί"
    rp2.font.size = Pt(16)
    rp2.font.color.rgb = DARK

    # Arrow between
    arr = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW, Inches(6.25), Inches(2.35), Inches(0.5), Inches(0.4)
    )
    arr.fill.solid()
    arr.fill.fore_color.rgb = MID_GRAY
    arr.line.fill.background()


def slide_stress_test(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, LIGHT_GRAY)
    add_header_bar(slide, "Stress Test — 3 Ερωτήσεις", "Ρώτα πάντα όταν βλέπεις πλατφόρμα ή broker", RGBColor(0xC5, 0x30, 0x30))

    questions = [
        (
            "1",
            "Αν αυτός ο intermediate εξαφανιστεί:",
            [
                "Το trade ακυρώνεται;",
                "Εκτελείται κανονικά μέσω CCP;",
                'Ή «μένει στον αέρα»;',
            ],
        ),
        (
            "2",
            "Πού είναι το legal counterparty;",
            ["Broker;", "CCP;", "Τελικός client;"],
        ),
        (
            "3",
            "Πού κρατιούνται τα assets;",
            [
                "Segregated custodian;",
                "Omnibus account;",
                "Broker balance sheet;",
            ],
        ),
    ]

    y = 1.5
    for num, q, opts in questions:
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(y), Inches(12.1), Inches(1.45)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)

        badge = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(0.85), Inches(y + 0.25), Inches(0.55), Inches(0.55)
        )
        badge.fill.solid()
        badge.fill.fore_color.rgb = RISK_RED
        badge.line.fill.background()
        btf = badge.text_frame
        btf.vertical_anchor = MSO_ANCHOR.MIDDLE
        bp = btf.paragraphs[0]
        bp.text = num
        bp.font.size = Pt(18)
        bp.font.bold = True
        bp.font.color.rgb = WHITE
        bp.alignment = PP_ALIGN.CENTER

        qbox = slide.shapes.add_textbox(Inches(1.6), Inches(y + 0.15), Inches(10.8), Inches(1.2))
        qtf = qbox.text_frame
        qtf.word_wrap = True
        qp = qtf.paragraphs[0]
        qp.text = q
        qp.font.size = Pt(15)
        qp.font.bold = True
        qp.font.color.rgb = DARK
        qp.space_after = Pt(6)
        for opt in opts:
            op = qtf.add_paragraph()
            op.text = f"   → {opt}"
            op.font.size = Pt(12)
            op.font.color.rgb = MID_GRAY

        y += 1.65

    footer = slide.shapes.add_textbox(Inches(0.6), Inches(6.7), Inches(12.1), Inches(0.5))
    fp = footer.text_frame.paragraphs[0]
    fp.text = "Επόμενο βήμα: margining · dealer exposure παρά το DVP · Lehman-style failure chains"
    fp.font.size = Pt(11)
    fp.font.italic = True
    fp.font.color.rgb = MID_GRAY
    fp.alignment = PP_ALIGN.CENTER


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide_title(prs)
    slide_overview(prs)

    slide_layer(
        prs, 1, "Trading Layer", "π.χ. Tradeweb",
        [
            "Buyer και Seller συμφωνούν τιμή",
            'Το trade "matches"',
            "Δημιουργείται confirmation",
        ],
        [
            "Market risk (microseconds/minutes) μέχρι lock",
            "Execution risk (λάθος price / partial fills)",
        ],
        "👉 Δεν υπάρχει ακόμα transfer χρημάτων ή τίτλων",
    )

    slide_layer(
        prs, 2, "Intermediation Layer", "π.χ. Tradition",
        [
            "Φέρνει counterparties μαζί (voice ή electronic)",
            "Κρατά «νομική μεσολάβηση» στο trade",
            "Pure broker (agency) ή matched principal",
        ],
        [
            "Broker credit risk — πτώχευση πριν settlement, λάθος segregation",
            "Operational risk — λάθη trade details, mismatch confirmations",
        ],
        "👉 Μπορεί να χαθεί operational control του trade",
        extra_sections=[
            ("Τύποι intermediation:", ["Agency broker (pure broker)", "Matched principal (προσωρινό exposure)"]),
        ],
    )

    slide_layer(
        prs, 3, "Clearing & Settlement", "DVP / CCP",
        [
            "DVP: τίτλος ↔ cash ανταλλάσσονται ταυτόχρονα",
            "CCP (αν υπάρχει): γίνεται ενδιάμεσος counterparty",
        ],
        [
            "Principal risk — ΣΧΕΔΟΝ εξαφανίζεται (δεν πληρώνεις χωρίς delivery)",
            "CCP default risk — σπάνιο αλλά systemic",
            "Settlement fail risk — technical/operational, fails to deliver",
        ],
        "👉 Το κρίσιμο σημείο του συστήματος",
    )

    slide_layer(
        prs, 4, "Custody Layer", "Τελική αποθήκη assets",
        [
            "Assets σε custodian bank",
            "Omnibus accounts ή segregated accounts",
        ],
        [
            "Custodian insolvency risk — τι γίνεται με segregation;",
            "Legal ownership risk — «whose name is the asset in?»",
        ],
        None,
    )

    slide_flow(prs)
    slide_insight(prs)
    slide_stress_test(prs)

    out = "/workspace/Trading_System_Risk_Layers.pptx"
    prs.save(out)
    print(f"Saved: {out}")


if __name__ == "__main__":
    main()
